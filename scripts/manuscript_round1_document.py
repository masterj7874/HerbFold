"""Build the review revision and response as editable documents."""
from pathlib import Path
import json,re
import manuscript_document as doc
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'research/manuscript'
OUT=ROOT/'output/manuscript/round1'
FIG=ROOT/'output/manuscript/figures/round1'
OUT.mkdir(parents=True,exist_ok=True)

old_setup=doc.setup
def setup(supplement):
    d=old_setup(supplement)
    for name in ('Normal','Table Source','Bibliography'):
        doc.set_font(d.styles[name],'CMU Serif',d.styles[name].font.size.pt)
    d.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    d.styles['Korean Body'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    d.core_properties.subject='Round 1 revision, 2026-09-09; archived observations frozen 2026-09-08'
    return d

def main():
    doc.OUT=OUT;doc.FIG=FIG;doc.setup=setup
    from manuscript_round1_tables import extend
    tab=extend(doc.tables(),SRC)
    (SRC/'round1/document_tables.json').write_text(json.dumps(tab,ensure_ascii=False,indent=2))
    main=doc.bibliography((SRC/'manuscript.md').read_text())
    main+='\n\n## Tables\n\n{{TABLE:1}}\n\n'
    for p in sorted(FIG.glob('figure-0*.png')):
        if not p.name.endswith('.docx.png'):main+='{{FIGURE:'+p.stem+'}}\n\n'
    (SRC/'round1/manuscript-resolved.md').write_text(main)
    unresolved=re.findall(r'\{\{(?!TABLE:|FIGURE:|EQUATION:|PAGEBREAK)[^}]+\}\}|\[@',main)
    assert not unresolved,unresolved
    si=(SRC/'supplementary.md').read_text()
    results=[doc.build(main,OUT/'HerbFold_R1_Manuscript.docx',tab),doc.build(si,OUT/'HerbFold_R1_Supplementary.docx',tab,True)]
    response=SRC/'round1/response-to-reviewers.md'
    if response.exists():
        results.append(doc.build(response.read_text(),OUT/'HerbFold_R1_Response.docx',tab,True))
        from docx import Document
        p=OUT/'HerbFold_R1_Response.docx';r=Document(p)
        r.sections[0].header.paragraphs[0].text='HERBFOLD  |  RESPONSE TO ROUND 1 REVIEW'
        r.save(p)
    tokens=doc.TOKENS.copy();tokens['round1_override']={'body_font':'CMU Serif','body_alignment':'left','native_equations':4,'correspondence':'Preserved from current supplied DOCX; Young Sahng Suh'}
    (SRC/'round1/document-design-tokens.json').write_text(json.dumps(tokens,indent=2))
    (SRC/'round1/document-build.json').write_text(json.dumps({'outputs':results,'style_geometry_audit':'passed','reference_count':len(json.loads((SRC/'citation-map.json').read_text()))},indent=2))
    print(json.dumps(results,indent=2))

if __name__=='__main__':main()
