"""Build the research manuscript and supplement from frozen evidence, no inference.

Run with the managed document Python runtime. Scientific computations and figure
generation are separate. This builder changes only manuscript outputs.
"""
from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / 'research/manuscript'
OUT = ROOT / 'output/manuscript'
FIG = OUT / 'figures' / 'scientific-reports'
OUT.mkdir(parents=True, exist_ok=True)

TOKENS = {
    'base_preset': 'narrative_proposal',
    'header_pattern': 'editorial title with quiet running label',
    'named_override': 'academic_reference_manuscript',
    'override_reason': 'Supplied journal articles: serif research prose, compact black hierarchy, original multipanel figures; no publisher branding.',
    'page': {'width_in': 8.5, 'height_in': 11, 'margin_in': 1, 'header_in': .492, 'footer_in': .492, 'usable_dxa': 9360},
    'body': {'font': 'Liberation Serif', 'size_pt': 11, 'before_pt': 0, 'after_pt': 6, 'line_multiple': 1.10, 'alignment': 'justified'},
    'title': {'font': 'Liberation Sans', 'size_pt': 22, 'before_pt': 0, 'after_pt': 12, 'color': '183B56', 'bold': True},
    'h1': {'font': 'Liberation Sans', 'size_pt': 14, 'before_pt': 16, 'after_pt': 8, 'color': '183B56'},
    'h2': {'font': 'Liberation Sans', 'size_pt': 11.5, 'before_pt': 12, 'after_pt': 6, 'color': '202D38'},
    'h3': {'font': 'Liberation Sans', 'size_pt': 11, 'before_pt': 8, 'after_pt': 4, 'color': '202D38'},
    'table': {'width_dxa': 9360, 'indent_dxa': 120, 'cell_margins_dxa': {'top': 80, 'bottom': 80, 'start': 120, 'end': 120}, 'font_pt': 9, 'header_fill': 'EAF0F3', 'border': 'D1DAE0'},
    'table_source': {'before_pt': 4, 'after_pt': 4, 'size_pt': 9},
    'caption': {'size_pt': 9.5, 'line_multiple': 1.08, 'after_pt': 8},
    'bibliography': {'size_pt': 9.5, 'line_multiple': 1.05, 'after_pt': 5, 'indent_dxa': 440, 'hanging_dxa': 440},
    'korean_override': {'east_asia_font': 'NanumGothic', 'body_size_pt': 10.5, 'line_multiple': 1.15},
    'journal_override': {'target': 'Scientific Reports', 'main_display_items': 8, 'main_line_numbers_every': 5, 'document_figure_dpi': 240, 'standalone_figure_dpi': 600},
    'figure_override': {'width_in': 6.5, 'max_height_in': 6.6, 'figure_starts_new_page': True},
    'author_metadata_override': {'font': 'Liberation Sans', 'author_size_pt': 11, 'affiliation_size_pt': 9.5, 'line_multiple': 1.05, 'main_title_page': True, 'source': 'HANBIT_2026-09-01.docx'},
}


def rows(name):
    with (SRC / 'tables' / name).open() as f:
        return list(csv.DictReader(f))


def num(v, places=3):
    return f'{float(v):.{places}f}'


def tables():
    t = {}
    def add(key, title, columns, values, widths, note, size=9):
        assert sum(widths) == 9360 and all(len(x) == len(columns) for x in values)
        t[key] = dict(title=title, columns=columns, rows=values, widths_dxa=widths, note=note, size=size)
    add('1', 'Table 1. Separate evidence cohorts and denominators',
        ['Analysis', 'Count', 'Unit and supported interpretation'], [
        ['Local catalog', '766,417', 'Distinct stored structures; not authenticated herbal materials'],
        ['Botanical retrieval', '1,293', 'Structures across four queried herb groups'],
        ['Enumeration', '5,000,000', 'Attempted fragment combinations'],
        ['Retained generation', '3,701,740', 'Unique structures after recorded chemical/lineage checks'],
        ['Candidate assay assessment', '15,740', 'Prior candidates plus a sampled generation subset'],
        ['AF3 comparison', '2 ligands; 1 target', 'Four jobs; 20 diffusion samples from one seed per job'],
        ['Quantum pilot', '4 inputs', 'Descriptor vectors; 156 physical qubits'],
        ['Tox21 model data', '7,617', 'Curated structures with endpoint-specific missing labels'],
        ['Qualified hERG assay', '150', 'Structures in a radioligand competition assay'],
    ], [2500, 1720, 5140], 'Counts are not additive. Biological assessment does not cover all generated structures. Source: numerical ledger and source-data CSVs.')
    ar = rows('af3_top_copies.csv')
    lookup = {(r['molecule'], r['condition']): r for r in ar}
    order = [('aspirin', 'baseline'), ('aspirin', 'msa'), ('quercetin', 'baseline'), ('quercetin', 'msa')]
    tr = []
    for label, key, n in [('pTM', 'ptm', 2), ('ipTM', 'iptm', 2), ('Protein C-alpha pLDDT', 'protein_ca_mean_plddt', 2), ('Ligand heavy-atom pLDDT', 'ligand_heavy_mean_plddt', 2), ('Protein RMSD to 5IKR (Å)', 'reference_ca_rmsd_angstrom', 3)]:
        tr.append([label] + [num(lookup[k][key], n) for k in order])
    add('2', 'Table 2. Top-ranked AF3 outputs for two ligand-target pairs',
        ['Metric', 'Aspirin\nbaseline', 'Aspirin\nstandard', 'Quercetin\nbaseline', 'Quercetin\nstandard'], tr,
        [3040, 1580, 1580, 1580, 1580], 'Baseline: no MSA or templates. Standard: MSA plus templates. Each job: seed 1, five diffusion samples, ten recycles. RMSD uses 551 matched C-alpha atoms. The reference contains mefenamic acid; no ligand-pose accuracy is inferred. Source: af3_top_copies.csv.')
    g = rows('gqd_taxon_source_counts.csv'); gl = {(r['herb_id'], r['source_id']): r for r in g}
    names = [('gegen', 'Gegen / Pueraria group*'), ('huangqin', 'Huangqin / S. baicalensis'), ('huanglian', 'Huanglian / C. chinensis'), ('gancao', 'Gancao / G. uralensis')]
    vals = [[name] + [gl[(k, src)]['unique_structures'] for src in ['lotus-2026-04', 'coconut-2026-09', 'union']] + [gl[(k, 'union')]['provenance_records']] for k, name in names]
    vals.append(['Four-group union', '590', '1,171', '1,293', '2,744†'])
    add('3', 'Table 3. Exact-taxon retrieval specialized to four medicinal herbs',
        ['Herb and recorded taxon', 'LOTUS\nstructures', 'COCONUT\nstructures', 'Union\nstructures', 'Herb-record\nmemberships'], vals,
        [3440, 1400, 1560, 1400, 1560], '* Pueraria montana var. lobata plus explicitly included synonym P. lobata. † 2,744 herb-record memberships represent 2,687 distinct provenance IDs. Source counts overlap; union values are deduplicated structures. No preparation, concentration or material authentication is implied. Source: gqd_taxon_source_counts.csv and gqd_retrieval.json.')
    execrows = []
    for k in order:
        r = lookup[k];execrows.append([k[0].capitalize(), 'No MSA/templates' if k[1]=='baseline' else 'MSA + templates', r['job_id']])
    execrows += [['Global quantum', '10 circuits × 128 shots', 'dafn10dnj4cs73agj6fg'], ['Projected quantum', '17 circuits × 1,024 shots', 'dafncudnj4cs73agjjsg']]
    add('S1', 'Supplementary Table S1. Execution identifiers retained for the paper', ['Branch / ligand', 'Condition', 'Local AF3 job or IBM provider job ID'], execrows, [1800, 2660, 4900], 'AF3 entries identify local jobs. Quantum entries identify IBM provider jobs. Projected local store job: f1b0f1c6867b4d21b40193e65ab980bf. Original quantum source analysis: b3d261566fa14891925d3b1fd730f2d1. Source hashes appear in artifact_hashes.csv.', 8.7)
    add('S2', 'Supplementary Table S2. Complementary roles of the three supplied studies',
        ['Study', 'Primary experimental unit', 'Evidence used as motivation'], [
        ['Xu et al. (2015)', 'GQD clinical intervention; 224 randomized, 187 analyzed', 'Formula context and microbiota-linked observations'],
        ['Abramson et al. (2024)', 'Biomolecular structure prediction benchmarks', 'AF3 coordinates and confidence interpretation'],
        ['Ghazi Vakili et al. (2025)', '16-qubit generative component; 15 compounds synthesized', 'Separation of quantum generation from downstream assays'],
        ['HerbFold (this work)', 'Separate chemistry, AF3, QPU and assay cohorts', 'Traceable evidence and conditional inference'],
    ], [2000, 3750, 3610], 'This is a design comparison, not a head-to-head performance benchmark. Clinical observations and synthesis in the cited studies were not performed here. Full references are in the main paper.')
    tox = rows('tox21_endpoint_metrics.csv')
    add('S3A', 'Supplementary Table S3a. All Tox21 endpoint performance estimates',
        ['Endpoint', 'Train / test n', 'Test positives', 'ROC-AUC\n[95% range]', 'AP', 'Test prevalence'],
        [[r['endpoint'], f"{r['train_count']} / {r['test_count']}", r['test_positive'], f"{num(r['roc_auc'])}\n[{num(r['roc_auc_lower95'])}, {num(r['roc_auc_upper95'])}]", num(r['average_precision']), num(r['prevalence'])] for r in tox],
        [1870, 1510, 1210, 2200, 1130, 1440], 'ROC-AUC ranges are stored scaffold-bootstrap percentiles. AP: average precision. Missing labels yield different endpoint sample counts. All twelve implemented reporting gates passed on these retrospective data; the gate is not independently prospectively validated. Source: tox21_endpoint_metrics.csv.', 8.5)
    add('S3B', 'Supplementary Table S3b. Tox21 loss and applicability coverage',
        ['Endpoint', 'Brier loss', 'Baseline Brier', 'In-domain test n', 'Candidate scores n', 'No inferred score n'],
        [[r['endpoint'], num(r['brier'],4), num(r['train_prevalence_brier_baseline'],4), r['test_in_domain'], r['candidate_predictions'], r['candidate_abstentions']] for r in tox],
        [1870, 1260, 1420, 1510, 1690, 1610], 'Baseline predicts the training prevalence for every test molecule. Candidate denominator: 15,740 per endpoint. No-score counts can include observed-assay precedence and are not equivalent to an out-of-domain count. The scores are uncalibrated assay estimates. Source: tox21_endpoint_metrics.csv.', 8.5)
    a = rows('af3_samples.csv')
    add('S4', 'Supplementary Table S4. All twenty AF3 diffusion samples',
        ['Ligand', 'Condition', 'ID', 'pTM', 'ipTM', 'Protein\npLDDT', 'Ligand\npLDDT', 'RMSD\n(Å)'],
        [[r['molecule'].capitalize(), 'Baseline' if r['condition']=='baseline' else 'Standard', r['sample'], num(r['ptm'],2), num(r['iptm'],2), num(r['protein_ca_mean_plddt'],2), num(r['ligand_heavy_mean_plddt'],2), num(r['reference_ca_rmsd_angstrom'],3)] for r in a],
        [1450, 1420, 740, 780, 780, 1390, 1390, 1410], 'Every row uses seed 1; sample IDs are local to their job. Protein pLDDT is C-alpha-only, ligand pLDDT heavy-atom-only. Standard adds MSA and templates together. Four top-file copies are excluded from this 20-row sample count. Source: af3_samples.csv.', 8.6)
    d = json.loads((SRC / 'data_audit.json').read_text());q = d['sections']['quantum']
    labels = ['C1', 'C2', 'Q', 'A']
    add('S5', 'Supplementary Table S5. Measured projected quantum kernel', ['Input'] + labels, [[labels[i]] + [num(v,6) for v in q['kernel'][i]] for i in range(4)], [1560, 1950, 1950, 1950, 1950], 'C1/C2 are the recorded BRICS candidates; Q: quercetin; A: aspirin. Diagonal values are algebraic. This is a descriptor-similarity matrix, not a binding affinity matrix. Exact values, classical/ideal references and raw resampling ranges are in quantum_pair_metrics.csv.')
    h = rows('herg_metrics.csv');h = [r for r in h if r['partition'] in ['test', 'test_in_domain']]
    add('S6', 'Supplementary Table S6. hERG binding model and baselines', ['Model', 'Test subset', 'n', 'MAE', 'RMSE', 'R²'],
        [[r['model'].replace('_',' '), 'All' if r['partition']=='test' else 'In domain', r['n'], num(r['mae']), num(r['rmse']), num(r['r2'])] for r in h], [2630, 1880, 750, 1350, 1350, 1400], 'MAE/RMSE use pIC50 units for CHEMBL1827362. Random forest was selected using tuning data; qualification also used held-out criteria. In-domain interval coverage: 21/23 = 0.913; half-width 0.647691 pIC50. This competitive radioligand assay is not a cardiac safety test. Source: herg_metrics.csv and evidence ledger.')
    from manuscript_q1_tables import extend_tables
    t=extend_tables(t,SRC)
    (SRC / 'document_tables.json').write_text(json.dumps(t, ensure_ascii=False, indent=2))
    return t


def bibliography(text):
    refs = json.loads((SRC / 'references_candidates.json').read_text())['references']
    lookup = {r['key']: r for r in refs}; order = []
    def cite(m):
        keys = [k.strip().lstrip('@') for k in m.group(1).split(';')]
        for k in keys:
            if k not in lookup: raise ValueError(f'Unknown citation {k}')
            if k not in order: order.append(k)
        return '[' + ', '.join(str(order.index(k)+1) for k in keys) + ']'
    text = re.sub(r'\[(@[^\]]+)\]', cite, text)
    lines = [];bib = []
    for n, key in enumerate(order, 1):
        r = lookup[key];authors = r.get('authors', [])
        details = r.get('author_details', [])
        names = [x['family'] + ', ' + ''.join(w[0]+'.' for w in x.get('given','').replace('-',' ').split() if w) for x in details if x.get('family')]
        names = names or authors
        a = (names[0] + ' et al.') if len(names)>5 else ', '.join(names)
        a = a or r.get('organization', 'Official documentation')
        year = f"({r['year']})" if r.get('year') else f"(accessed {r.get('accessed', '2026-09-08')})"
        journal = r.get('journal', '')
        vol = r.get('volume','');pages = r.get('pages') or r.get('article_number','')
        journal_part = ' '.join(str(v) for v in [journal, vol, pages] if v)
        url = r.get('url', r.get('source_url',''))
        line = f"[{n}] {a.rstrip('.')}. {r['title']}. {journal_part} {year}. {url}"
        lines.append(re.sub(r' +',' ',line))
        fields = {'title': r['title'], 'author': ' and '.join(authors), 'year': str(r.get('year') or ''), 'journal': journal, 'volume': vol, 'pages': pages, 'doi': r.get('doi',''), 'url': url}
        bib.append('@misc{' + key + ',\n' + ',\n'.join(f'  {k} = {{{str(v)}}}' for k,v in fields.items() if v) + '\n}')
    text = text.replace('{{REFERENCES}}','\n\n'.join(lines))
    (SRC / 'references.bib').write_text('\n\n'.join(bib)+'\n')
    (SRC / 'citation-map.json').write_text(json.dumps({key:n for n,key in enumerate(order,1)}, indent=2))
    return text


def el(tag, **attrs):
    e = OxmlElement('w:'+tag)
    for k,v in attrs.items(): e.set(qn('w:'+k), str(v))
    return e


def set_font(style, name, size, color='202020', bold=False):
    style.font.name = name;style.font.size = Pt(size);style.font.color.rgb = RGBColor.from_string(color);style.font.bold=bold;style.font.italic=False
    rpr = style.element.get_or_add_rPr();rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf=el('rFonts');rpr.append(rf)
    for key,val in {'ascii':name,'hAnsi':name,'eastAsia':'NanumGothic','cs':name}.items():rf.set(qn('w:'+key),val)


def style(doc, name, font, size, before=0, after=0, line=1.1, color='202020', bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, keep=False):
    s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    set_font(s,font,size,color,bold);p=s.paragraph_format
    p.space_before=Pt(before);p.space_after=Pt(after);p.line_spacing=line;p.alignment=align;p.widow_control=True;p.keep_with_next=keep
    for old in list(s.element.get_or_add_pPr().findall(qn('w:pBdr'))):s.element.get_or_add_pPr().remove(old)
    return s


def setup(supplement):
    doc = Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=sec.footer_distance=Inches(.492)
    if not supplement:sec._sectPr.append(el('lnNumType',countBy=5,start=1,restart='continuous',distance=180))
    style(doc,'Normal','Liberation Serif',11,after=6,line=1.10,align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    style(doc,'Title','Liberation Sans',22,after=12,color='183B56',bold=True,keep=True)
    style(doc,'Subtitle','Liberation Sans',10,after=8,color='596873',keep=True)
    style(doc,'Authors','Liberation Sans',11,after=10,line=1.10,bold=True,keep=True)
    style(doc,'Affiliation','Liberation Sans',9.5,after=5,line=1.05)
    style(doc,'Author Information','Liberation Sans',9.5,before=5,after=7,line=1.05)
    for n,size,bef,aft,col in [(1,14,16,8,'183B56'),(2,11.5,12,6,'202D38'),(3,11,8,4,'202D38')]:
        style(doc,f'Heading {n}','Liberation Sans',size,before=bef,after=aft,color=col,bold=True,keep=True)
    style(doc,'Caption','Liberation Sans',9.5,after=8,line=1.08)
    style(doc,'Table Caption','Liberation Sans',10,before=12,after=7,line=1.08,bold=True,keep=True)
    style(doc,'Table Text','Liberation Sans',9,line=1.04)
    style(doc,'Table Source','Liberation Serif',9,before=4,after=4,line=1.04)
    style(doc,'Bibliography','Liberation Serif',9.5,after=5,line=1.05)
    style(doc,'Korean Body','NanumGothic',10.5,after=6,line=1.15,align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    style(doc,'Figure Image','Liberation Sans',1,after=8,line=1,keep=True)
    style(doc,'Header','Liberation Sans',8,line=1,color='667681')
    style(doc,'Footer','Liberation Sans',8,line=1,color='667681',align=WD_ALIGN_PARAGRAPH.RIGHT)
    h=sec.header.paragraphs[0];h.style='Header';h.text='HERBFOLD  |  '+('SUPPLEMENTARY INFORMATION' if supplement else 'RESEARCH ARTICLE')
    foot=sec.footer.paragraphs[0];foot.style='Footer';foot.add_run('HerbFold  •  ')
    field=el('fldSimple',instr='PAGE');foot._p.append(field)
    abstract=el('abstractNum',abstractNumId=77);abstract.append(el('multiLevelType',val='singleLevel'))
    lvl=el('lvl',ilvl=0);lvl.append(el('start',val=1));lvl.append(el('numFmt',val='decimal'));lvl.append(el('lvlText',val='[%1]'));lvl.append(el('lvlJc',val='left'))
    pp=el('pPr');pp.append(el('ind',left=440,hanging=440));tabs=el('tabs');tabs.append(el('tab',val='num',pos=440));pp.append(tabs);lvl.append(pp);abstract.append(lvl)
    doc.part.numbering_part.element.append(abstract)
    nn=el('num',numId=77);nn.append(el('abstractNumId',val=77));doc.part.numbering_part.element.append(nn)
    doc.core_properties.title='HerbFold: '+('Supplementary Information' if supplement else 'Auditable medicinal-plant discovery links chemical provenance with AlphaFold 3 diagnostics and quantum feature measurements')
    doc.core_properties.author='Young Sahng Suh; Hwa-Young Lee; Shuji Ogino'
    doc.core_properties.subject='Computational methods; actual archived results; data freeze 2026-09-08'
    return doc


def runs(p, text):
    # Keep visible text intact; references and formulas remain editable text.
    text = text.replace(' C-alpha',' C-α').replace(' angstroms',' Å')
    text = re.sub(r'(\d) A(?=[.,; ])',r'\1 Å',text)
    for i,part in enumerate(re.split(r'(\*\*.*?\*\*)',text)):
        r=p.add_run(part[2:-2] if part.startswith('**') else part)
        if part.startswith('**'):r.bold=True


def math_run(text, plain=False):
    r=OxmlElement('m:r')
    if plain:
        prop=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');prop.append(sty);r.append(prop)
    wp=OxmlElement('w:rPr');wp.append(el('sz',val=22));r.append(wp)
    t=OxmlElement('m:t');t.set(qn('xml:space'),'preserve');t.text=text;r.append(t);return r


def math_sub(base, sub, plain_sub=False):
    node=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(math_run(base));node.append(e)
    q=OxmlElement('m:sub');q.append(math_run(sub,plain_sub));node.append(q);return node


def math_sup(items, sup):
    node=OxmlElement('m:sSup');e=OxmlElement('m:e')
    for i in items:e.append(i)
    node.append(e);q=OxmlElement('m:sup');q.append(math_run(sup));node.append(q);return node


def math_fraction(numerator, denominator):
    f=OxmlElement('m:f')
    for tag,items in [('m:num',numerator),('m:den',denominator)]:
        e=OxmlElement(tag)
        for i in items:e.append(i)
        f.append(e)
    return f


def math_sum(index, expression):
    n=OxmlElement('m:nary');p=OxmlElement('m:naryPr')
    for name,value in [('chr','∑'),('limLoc','subSup'),('supHide','1')]:
        e=OxmlElement('m:'+name);e.set(qn('m:val'),value);p.append(e)
    n.append(p);s=OxmlElement('m:sub');s.append(math_run(index));n.append(s)
    n.append(OxmlElement('m:sup'));e=OxmlElement('m:e')
    for x in expression:e.append(x)
    n.append(e);return n


def equation(doc,number,expression):
    R=math_run;S=math_sub;F=math_fraction
    if number=='2':
        items=[S('D','ij'),R(' = '),F([R('1')],[R('2Q')]),
               math_sum('q',[math_sup([R('‖'),S('r','iq'),R(' − '),S('r','jq'),R('‖')],'2')]),
               R(',     '),S('K','ij'),R(' = '),R('exp',True),R('(−γ'),S('D','ij'),R(').')]
    elif number=='3':
        items=[S('K','c'),R(' = HKH,   H = I − '),F([R('1'),math_sup([R('1')],'T')],[R('n')]),
               R(',   A(K,L) = '),F([R('⟨'),S('K','c'),R(','),S('L','c'),S('⟩','F')],
               [R('‖'),S('K','c'),S('‖','F'),R(' ‖'),S('L','c'),S('‖','F')]),R('.')]
    elif number=='4':
        items=[S('r','eff',True),R(' = '),R('exp',True),R('('),R('−'),
               math_sum('j',[S('p','j'),R(' '),R('log',True),R(' '),S('p','j')]),R('),    '),S('p','j'),R(' = '),
               F([S('λ','j')],[math_sum('k',[S('λ','k')])]),R('.')]
    else:items=[R(expression)]
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(7);p.paragraph_format.space_after=Pt(7)
    o=OxmlElement('m:oMath')
    for item in items:o.append(item)
    p._p.append(o)
    # Word text preserves the label gap that the math renderer collapses.
    label=p.add_run('\u2003\u2003('+number+')');label.font.size=Pt(11)


def table(doc, spec):
    doc.add_paragraph(spec['title'],style='Table Caption')
    tb=doc.add_table(rows=1+len(spec['rows']),cols=len(spec['columns']));tb.autofit=False
    pr=tb._tbl.tblPr
    for tag in ['tblW','tblInd','tblLayout','tblCellMar','tblBorders']:
        for old in pr.findall(qn('w:'+tag)):pr.remove(old)
    pr.append(el('tblW',w=9360,type='dxa'));pr.append(el('tblInd',w=120,type='dxa'));pr.append(el('tblLayout',type='fixed'))
    mar=el('tblCellMar')
    for k,v in TOKENS['table']['cell_margins_dxa'].items():
        if k in {'top','bottom'}:v=spec.get('vertical_padding_dxa',v)
        mar.append(el(k,w=v,type='dxa'))
    pr.append(mar);borders=el('tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:borders.append(el(side,val='single',sz=4,color='D1DAE0'))
    pr.append(borders)
    grid=tb._tbl.tblGrid
    for child in list(grid):grid.remove(child)
    for w in spec['widths_dxa']:grid.append(el('gridCol',w=w))
    for ri,values in enumerate([spec['columns']]+spec['rows']):
        row=tb.rows[ri];row._tr.get_or_add_trPr().append(el('cantSplit'))
        if ri==0:row._tr.get_or_add_trPr().append(el('tblHeader'))
        for ci,value in enumerate(values):
            c=row.cells[ci];c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;c.width=Inches(spec['widths_dxa'][ci]/1440)
            cp=c._tc.get_or_add_tcPr();w=cp.find(qn('w:tcW'));w.set(qn('w:w'),str(spec['widths_dxa'][ci]));w.set(qn('w:type'),'dxa')
            if ri==0:cp.append(el('shd',fill='EAF0F3'))
            p=c.paragraphs[0];p.style='Table Text';p.paragraph_format.keep_with_next=True
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or len(str(value))>35 else WD_ALIGN_PARAGRAPH.CENTER
            rr=p.add_run(str(value));rr.font.size=Pt(spec['size']);rr.bold=ri==0
    doc.add_paragraph(spec['note'],style='Table Source')


def figure(doc, stem):
    p=doc.add_paragraph(style='Figure Image')
    p.paragraph_format.page_break_before=True
    im=FIG/(stem+'.docx.png')
    if not im.exists():im=FIG/(stem+'.png')
    with Image.open(im) as img:w,h=img.size
    width=min(6.5,6.6*w/h)
    p.add_run().add_picture(str(im),width=Inches(width));p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cap=(FIG/(stem+'.caption.txt')).read_text().strip()
    p=doc.add_paragraph(style='Caption')
    m=re.match(r'((?:Supplementary )?Figure [^.]+\.)(.*)',cap,flags=re.S)
    if m:p.add_run(m.group(1)).bold=True;p.add_run(m.group(2))
    else:p.add_run(cap)


def build(text, path, tab, supplement=False):
    doc=setup(supplement);korean=False;firstbody=True
    for block in re.split(r'\n\s*\n',text.strip()):
        b=block.strip()
        if not b:continue
        if b.startswith('{{EQUATION:'):
            number, expression=b.removeprefix('{{EQUATION:').removesuffix('}}').split('|',1)
            equation(doc,number,expression)
            continue
        if b.startswith('{{TABLE:'):
            key=b.removeprefix('{{TABLE:').removesuffix('}}')
            table(doc,tab[key]);continue
        if b.startswith('{{FIGURE:'):
            figure(doc,b.removeprefix('{{FIGURE:').removesuffix('}}'));continue
        if b=='{{PAGEBREAK}}':doc.add_page_break();continue
        if b.startswith('# '):doc.add_paragraph(b[2:],style='Title');continue
        if b.startswith('Authors: '):
            p=doc.add_paragraph(style='Authors')
            for part in re.split(r'(\[[0-9,]+\])',b.removeprefix('Authors: ')):
                rr=p.add_run(part[1:-1] if part.startswith('[') else part)
                if part.startswith('['):rr.font.superscript=True
            continue
        if match:=re.match(r'^Affiliation (\d+): (.*)',b):
            p=doc.add_paragraph(style='Affiliation');p.add_run(match[1]).font.superscript=True
            p.add_run(' '+match[2]);continue
        if b.startswith(('Correspondence: ', 'Author positions: ', 'Author information: ')):
            p=doc.add_paragraph(style='Author Information')
            label,body=b.split(': ',1);p.add_run(label+': ').bold=True;p.add_run(body);continue
        if b.startswith('## '):
            title=b[3:]
            if title=='국문 연구 요약':korean=True;doc.add_page_break()
            p=doc.add_paragraph(title,style='Heading 1')
            if title=='Abstract' and not supplement:p.paragraph_format.page_break_before=True
            continue
        if b.startswith('### '):doc.add_paragraph(b[4:],style='Heading 2');continue
        if b.startswith('[') and re.match(r'^\[\d+\] ',b):
            p=doc.add_paragraph(style='Bibliography');ppr=p._p.get_or_add_pPr();np=el('numPr');np.append(el('ilvl',val=0));np.append(el('numId',val=77));ppr.append(np)
            runs(p,re.sub(r'^\[\d+\] ','',b));continue
        if firstbody:
            doc.add_paragraph(b,style='Subtitle');firstbody=False;continue
        p=doc.add_paragraph(style='Korean Body' if korean else 'Normal');runs(p,b.replace('\n',' '))
    doc.save(path)
    # Explicit geometry audit on every table, independent of visual rendering.
    for tb in doc.tables:
        assert tb._tbl.tblPr.find(qn('w:tblW')).get(qn('w:w'))=='9360'
        widths=[int(c.get(qn('w:w'))) for c in tb._tbl.tblGrid]
        assert sum(widths)==9360
        for row in tb.rows:
            assert [int(c._tc.get_or_add_tcPr().find(qn('w:tcW')).get(qn('w:w'))) for c in row.cells]==widths
    return {'file':str(path.relative_to(ROOT)),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'figures':len(doc.inline_shapes)}


def main():
    tab=tables();text=bibliography((SRC/'manuscript.md').read_text())
    text+='\n\n## Tables\n\n{{TABLE:1}}\n'
    for p in sorted(p for p in FIG.glob('figure-0*.png') if not p.name.endswith('.docx.png')):text+='\n{{FIGURE:'+p.stem+'}}\n'
    # Separate figure directives into blocks for a stable paragraph parser.
    text=text.replace('}}\n{{','}}\n\n{{')
    (SRC/'manuscript-resolved.md').write_text(text)
    unresolved=re.findall(r'\{\{(?!TABLE:|FIGURE:|EQUATION:|PAGEBREAK)[^}]+\}\}|\[@',text)
    assert not unresolved,unresolved
    outputs=[build(text,OUT/'HerbFold_Manuscript.docx',tab),build((SRC/'supplementary.md').read_text(),OUT/'HerbFold_Supplementary.docx',tab,True)]
    (SRC/'document-design-tokens.json').write_text(json.dumps(TOKENS,indent=2))
    (SRC/'document-build.json').write_text(json.dumps({'outputs':outputs,'style_geometry_audit':'passed','reference_count':len(json.loads((SRC/'citation-map.json').read_text()))},indent=2))
    print(json.dumps(outputs,indent=2))


if __name__=='__main__':main()
