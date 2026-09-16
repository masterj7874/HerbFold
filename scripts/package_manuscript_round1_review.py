"""Verify and assemble the actual R1 review documents and nested source archive."""
from pathlib import Path
import argparse
import csv
import hashlib
import json
import re
import shutil
import zipfile
from datetime import datetime, timezone
from docx import Document
from lxml import etree
import fitz

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output/manuscript/round1"
R1 = ROOT / "research/manuscript/round1"


def sha(p):
    h = hashlib.sha256()
    with p.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


VERIFIER = '''"""Verify the local review packet inventory using only Python's standard library."""
from pathlib import Path
import hashlib,json,sys

root=Path(__file__).resolve().parent
manifest=json.loads((root/'REVIEW_PACKET_MANIFEST.json').read_text())
failures=[]
for row in manifest['files']:
    path=root/row['path']
    if not path.resolve().is_relative_to(root):
        failures.append({'path':row['path'],'reason':'unsafe path'});continue
    if not path.is_file():
        failures.append({'path':row['path'],'reason':'missing'});continue
    h=hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda:f.read(1024*1024),b''):h.update(block)
    if path.stat().st_size!=row['size_bytes'] or h.hexdigest()!=row['sha256']:
        failures.append({'path':row['path'],'reason':'size/hash mismatch'})
print(json.dumps({'passed':not failures,'checked_files':len(manifest['files']),'failures':failures},indent=2))
sys.exit(bool(failures))
'''

README = """# HerbFold R1 — complete local review packet

This packet supplies the actual revised manuscript, Supplementary Information,
point-by-point response and source archive together. It is a local review draft;
it has not been submitted, deposited publicly or approved by all named authors.

Start with manuscript/HerbFold_R1_Manuscript.pdf, then the Supplementary PDF and
response. Editable DOCX copies accompany all three. review/review-map.csv maps
40 display callouts (12 figures and 28 table panels) to files, fields, units and
replay commands. Machine-readable scientific inputs live in the nested
source/HerbFold_R1_Source_Package.zip; extract it into a separate directory and
start with PACKAGE_README.md and python verify_package.py.

Run python verify_review_packet.py after extracting this outer ZIP. Its manifest
checks the distributed bytes, not biological validity or provider authenticity.
For the numerical and software replay commands, use review/packet-replay-guide.md.
Both failed attempts and successful bounded CPU replays are retained. The later
wrapper automatically stages the distributed snapshots; no new AF3, IBM QPU or
platform LLM job is required for these selected replays. The full catalog,
generation database, complete MSA databases and AF3 parameters are excluded and
are distinguished from the included snapshots in excluded-inputs-and-terms.csv.

The source ZIP exceeds 50 MB. This outer ZIP is a local convenience bundle,
not a claim that a journal submission form accepts it as one attachment. Main
documents, SI and source data are separately available in the output directory.
No new blanket licence or third-party redistribution permission is inferred.

author-review/ contains the remaining author confirmations and a separate
AI-assisted internal diagnostic recheck. That recheck is not journal peer review,
not an official new R-Score and not a request to a reviewer to award a score.
The original supplied 57.5 score is preserved; 90-point attainment is not claimed.
Author roles, final approval, study-specific funding/conflicts and public release
rights remain open. No therapeutic efficacy, validated selected-ligand pose or
quantum advantage is asserted by this packet.

한국어 안내: 본문·보충자료·답변서를 함께 검토하는 완성된 로컬 묶음입니다.
원고만 제출해 근거 자료가 다시 빠지지 않도록 소스 ZIP과 보충자료를 함께
전달해야 합니다. 저자의 실제 승인·기여·연구비·이해상충 및 공개 배포 조건은
author-review/author-confirmations-ko.md에서 확인해야 합니다. 내부 진단 점수는
정식 재심사나 90점 달성 인증이 아닙니다.
"""


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--main-render", type=Path, default=Path("tmp/manuscript/r1-numbered-qa/main"))
    args = parser.parse_args()
    pdf_roots = {"Manuscript": ROOT / args.main_render,
                 "Supplementary": ROOT / "tmp/manuscript/r1-release-qa/supplement",
                 "Response": ROOT / "tmp/manuscript/r1-release-qa/response"}
    mapping = {}
    document_checks = []
    for label, directory in pdf_roots.items():
        name = "HerbFold_R1_" + label
        source_pdf = directory / (name + ".pdf")
        assert source_pdf.is_file(), source_pdf
        shutil.copy2(source_pdf, OUT / source_pdf.name)
        docx = OUT / (name + ".docx")
        doc = Document(docx)
        with fitz.open(source_pdf) as pdf:
            pages = len(pdf)
            pdf_text = "\n".join(page.get_text() for page in pdf)
        expected = {"Manuscript": (1, 7), "Supplementary": (27, 5), "Response": (0, 0)}[label]
        assert (len(doc.tables), len(doc.inline_shapes)) == expected
        with zipfile.ZipFile(docx) as z:
            xml = etree.fromstring(z.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
              "m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
        all_text = " ".join(xml.xpath("//w:t/text()|//m:t/text()", namespaces=ns))
        assert not re.search(r"\{\{(?:FIGURE|TABLE|EQUATION):|\[@", all_text)
        equations = len(xml.xpath("//m:oMath", namespaces=ns))
        assert equations == (4 if label == "Manuscript" else 0)
        assert pdf_text.strip()
        document_checks.append({"document": label, "pages": pages,
            "tables": len(doc.tables), "inline_figures": len(doc.inline_shapes),
            "native_equations": equations, "docx_sha256": sha(docx),
            "pdf_sha256": sha(source_pdf), "rendered_page_pngs": len(list(directory.glob("page-*.png")))})
        for suffix in ("docx", "pdf"):
            p = OUT / (name + "." + suffix)
            mapping["manuscript/" + p.name] = p

    source_zip = OUT / "HerbFold_R1_Source_Package.zip"
    with zipfile.ZipFile(source_zip) as z:
        assert z.testzip() is None
        manifest = json.loads(z.read("PACKAGE_MANIFEST.json"))
        for row in manifest["files"]:
            content = z.read(row["path"])
            assert len(content) == row["size_bytes"]
            assert hashlib.sha256(content).hexdigest() == row["sha256"]
        replay = json.loads((R1 / "replay-report.json").read_text())
        frozen = replay["algorithm_hashes"] + replay["scientific_input_hashes"]
        for row in frozen:
            assert hashlib.sha256(z.read(row["path"])).hexdigest() == row["sha256"], row["path"]
        with (R1 / "review-map.csv").open(newline="") as f:
            callouts = list(csv.DictReader(f))
        assert len(callouts) == 40
        evidence_paths = set()
        for row in callouts:
            for path in row["source_files"].split(" | "):
                assert path in z.namelist(), (row["callout"], path)
                evidence_paths.add(path)
        assert z.read("research/manuscript/round1/review-map.csv") == (R1 / "review-map.csv").read_bytes()
        for audit_file in (R1 / "reassessment").glob("*"):
            if audit_file.is_file() and audit_file.suffix in {".md", ".json"}:
                assert z.read(str(audit_file.relative_to(ROOT))) == audit_file.read_bytes(), audit_file
        # New helper replay is recorded against an earlier editorial package;
        # carry it forward only if its exact script/input hash lists still match.
        additional = json.loads((R1 / "final-source-replay.json").read_text())
        helper_rows = additional["combined_carry_forward_hashes"]
        assert len(helper_rows) == 354
        for row in helper_rows:
            assert hashlib.sha256(z.read(row["path"])).hexdigest() == row["sha256"], row["path"]
    text = (ROOT / "research/manuscript/manuscript.md").read_text()
    abstract = text.split("## Abstract", 1)[1].split("## Introduction", 1)[0]
    core = text.split("## Introduction", 1)[1].split("## Methods", 1)[0]
    refs = json.loads((ROOT / "research/manuscript/references_candidates.json").read_text())
    cited = {key.strip().lstrip("@") for group in re.findall(r"\[@([^\]]+)\]", text) for key in group.split(";")}
    reference_keys = {r["key"] for r in refs["references"]}
    assert cited <= reference_keys
    assert len(abstract.split()) <= 200 and len(core.split()) <= 4500
    assert len(refs["references"]) == refs["reference_count"] == 36
    caption_checks = []
    for p in sorted((ROOT / "output/manuscript/figures/round1").glob("*.caption.txt")):
        words = len(p.read_text().split())
        assert words <= 350, p
        caption_checks.append({"file": p.name, "words": words})
    assert len(caption_checks) == 12
    final = {"status": "passed", "created_at": datetime.now(timezone.utc).isoformat(),
        "scope": "Artifact integrity, document structure and scientific replay hash carry-forward; visual and numerical audits are separately supplied",
        "documents": document_checks, "reference_count": 36,
        "abstract_words_whitespace": len(abstract.split()), "core_words_excluding_methods_whitespace": len(core.split()),
        "mapped_callouts": len(callouts), "mapped_evidence_files": len(evidence_paths),
        "figure_captions": caption_checks, "source_manifest_files": len(manifest["files"]),
        "source_zip_sha256": sha(source_zip), "source_zip_size_bytes": source_zip.stat().st_size,
        "frozen_scientific_paths_carried_forward": len(frozen),
        "combined_scientific_and_helper_paths_carried_forward": len(helper_rows),
        "review_map_in_source_zip_exact": True,
        "author_approval_and_public_release_rights_confirmed": False,
        "journal_submission_performed": False, "ninety_point_score_claimed": False}
    receipt = OUT / "final-deliverable-verification.json"
    receipt.write_text(json.dumps(final, indent=2, ensure_ascii=False) + "\n")
    mapping["review/" + receipt.name] = receipt
    for name in ("review-map.csv", "packet-replay-guide.md", "excluded-inputs-and-terms.csv",
                 "replay-report.json", "replay-report.md", "final-source-replay.json", "final-source-replay.md"):
        mapping["review/" + name] = R1 / name
    for name in ("author-confirmations-ko.md",):
        mapping["author-review/" + name] = R1 / name
    for p in sorted((R1 / "reassessment").glob("*")):
        if p.is_file() and p.suffix in {".md", ".json"}:
            mapping["author-review/reassessment/" + p.name] = p
    for p in sorted((R1 / "qa").glob("*")):
        if p.is_file() and p.suffix in {".md", ".json"}:
            mapping["review/visual-qa/" + p.name] = p
    mapping["source/" + source_zip.name] = source_zip
    mapping["source/HerbFold_R1_Source_Package.verification.json"] = OUT / "HerbFold_R1_Source_Package.verification.json"
    generated = {"README.md": README.encode(), "verify_review_packet.py": VERIFIER.encode()}
    rows = [{"path": name, "sha256": sha(p), "size_bytes": p.stat().st_size}
            for name, p in sorted(mapping.items())]
    rows += [{"path": name, "sha256": hashlib.sha256(data).hexdigest(), "size_bytes": len(data)}
             for name, data in generated.items()]
    outer = {"schema_version": 1, "created_at": datetime.now(timezone.utc).isoformat(),
             "files": sorted(rows, key=lambda r: r["path"]),
             "manifest_self_hash_excluded": True, "local_review_only": True,
             "source_zip_has_its_own_manifest_and_verifier": True}
    package = OUT / "HerbFold_R1_Review_Packet.zip"
    with zipfile.ZipFile(package, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as z:
        for name, p in sorted(mapping.items()):
            z.write(p, name, compress_type=zipfile.ZIP_STORED if p.suffix == ".zip" else zipfile.ZIP_DEFLATED)
        for name, data in generated.items():
            z.writestr(name, data)
        z.writestr("REVIEW_PACKET_MANIFEST.json", json.dumps(outer, indent=2, ensure_ascii=False) + "\n")
    with zipfile.ZipFile(package) as z:
        assert z.testzip() is None
        for row in rows:
            content = z.read(row["path"])
            assert len(content) == row["size_bytes"] and hashlib.sha256(content).hexdigest() == row["sha256"]
    outer_receipt = {"passed": True, "archive": str(package.relative_to(ROOT)),
                     "sha256": sha(package), "size_bytes": package.stat().st_size,
                     "manifest_files": len(rows), "zip_entries": len(rows) + 1,
                     "all_six_document_files_included": True, "nested_source_zip_included": True}
    (OUT / "HerbFold_R1_Review_Packet.verification.json").write_text(json.dumps(outer_receipt, indent=2) + "\n")
    (OUT / "README_R1_KO.md").write_text(README)
    print(json.dumps({"verification": final, "outer_packet": outer_receipt}, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
